from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

BG=RGBColor(0x1B,0x14,0x10)
CREAM=RGBColor(0xF3,0xE8,0xD3)
GOLD=RGBColor(0xC9,0xA8,0x6A)
MUTED=RGBColor(0xB9,0xA7,0x8D)
DARK=RGBColor(0x2A,0x1F,0x18)
RED=RGBColor(0x8B,0x3A,0x2F)
TITLE_FONT="Songti SC"
BODY_FONT="PingFang SC"

prs=Presentation()
prs.slide_width=Inches(13.333)
prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def bg(slide,color=BG):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    s.fill.solid(); s.fill.fore_color.rgb=color; s.line.fill.background()
    s.shadow.inherit=False
    return s

def text(slide,x,y,w,h,s,size=22,color=CREAM,font=BODY_FONT,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=1.15):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    p.line_spacing=spacing
    r=p.add_run(); r.text=s
    r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return tb

def rule(slide,x,y,w,color=GOLD,h=0.02):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=color; s.line.fill.background(); s.shadow.inherit=False
    return s

def pic(slide,path,x,y,w,h):
    p=slide.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    # gold border
    b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    b.fill.background(); b.line.color.rgb=GOLD; b.line.width=Pt(1.2); b.shadow.inherit=False
    return p

def num(slide,x,y,n,color=GOLD):
    text(slide,x,y,0.8,0.5,n,size=15,color=color,font=BODY_FONT,bold=True)

def footer(slide,n):
    text(slide,11.9,7.06,1.0,0.25,f"{n:02d} / 06",size=10,color=MUTED,align=PP_ALIGN.RIGHT)

# 01 cover
s=prs.slides.add_slide(blank)
s.shapes.add_picture("_ppt_assets/cover_momo.jpg",0,0,prs.slide_width,prs.slide_height)
# left gradient overlay, keeps Momo visible on the right
s.shapes.add_picture("_ppt_assets/cover_gradient.png",0,0,prs.slide_width,prs.slide_height)
text(s,0.8,0.62,4.0,0.35,"ACC 2026  ·  DEMO",size=13,color=GOLD,bold=True)
text(s,0.8,1.55,7.0,1.2,"墨墨的洋馆",size=58,color=CREAM,font=TITLE_FONT,bold=False)
rule(s,0.85,2.85,1.3)
text(s,0.8,3.05,6.6,0.8,"一个会自己玩狼人杀的 AI 茶会",size=25,color=CREAM)
text(s,0.8,3.95,6.4,1.0,"桌宠、洋馆日常、八个各怀心事的人，\n和一局不会重样的茶会。",size=17,color=MUTED,spacing=1.5)
text(s,0.8,6.55,6.5,0.35,"Lucian 的工作室",size=12,color=MUTED)
text(s,10.6,6.55,2.2,0.35,"2026.10",size=12,color=MUTED,align=PP_ALIGN.RIGHT)

# 02 not a chatbox
s=prs.slides.add_slide(blank); bg(s)
num(s,0.82,0.55,"01")
text(s,1.25,0.48,7.0,0.75,"她不是聊天框",size=36,color=CREAM,font=TITLE_FONT)
rule(s,0.82,1.35,1.1)
text(s,0.85,1.75,5.9,0.55,"墨墨住在洋馆里。",size=22,color=CREAM)
text(s,0.85,2.35,5.9,0.9,"你学习、攒金币、过日子，她会吐槽，也会护短。",size=16,color=MUTED,spacing=1.45)
text(s,0.85,3.55,5.9,0.55,"第七天，你被请进一场茶会。",size=22,color=CREAM)
text(s,0.85,4.15,5.9,0.9,"没有墙，没有门，桌上坐着七个陌生人。",size=16,color=MUTED,spacing=1.45)
text(s,0.85,5.5,5.8,0.6,"“先别碰。太干净了。”",size=18,color=GOLD,font=TITLE_FONT)
text(s,0.85,6.05,5.8,0.35,"——墨墨",size=12,color=MUTED)
pic(s,"_ppt_assets/card_tearoom.jpg",7.25,0.95,5.25,5.9)
footer(s,2)

# 03 four steps
s=prs.slides.add_slide(blank); bg(s)
num(s,0.82,0.55,"02")
text(s,1.25,0.48,9.0,0.75,"一局茶会，四步",size=36,color=CREAM,font=TITLE_FONT)
rule(s,0.82,1.35,1.1)
steps=[
 ("01","入席","月见发牌。八个人，身份随机。"),
 ("02","发言","每个人都要点名、站边、说理由。"),
 ("03","投票","说投谁就投谁；改票要解释。"),
 ("04","收尾","狼赢、好人赢，或月见裁。"),
]
x0=0.82; bw=2.83; gap=0.27
for i,(n,h,d) in enumerate(steps):
    x=x0+i*(bw+gap)
    box=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(1.95),Inches(bw),Inches(3.7))
    box.fill.solid(); box.fill.fore_color.rgb=DARK; box.line.color.rgb=RGBColor(0x4A,0x3A,0x2A); box.line.width=Pt(0.8); box.shadow.inherit=False
    text(s,x+0.25,2.2,0.7,0.4,n,size=13,color=GOLD,bold=True)
    text(s,x+0.25,2.75,bw-0.5,0.6,h,size=26,color=CREAM,font=TITLE_FONT)
    text(s,x+0.25,3.55,bw-0.5,1.7,d,size=14,color=MUTED,spacing=1.4)
text(s,0.85,6.15,11.7,0.5,"三狼（含狼王）、预言家、女巫、猎人、村民。每局重抽。",size=14,color=MUTED)
footer(s,3)

# 04 they are not reading lines
s=prs.slides.add_slide(blank); bg(s)
num(s,0.82,0.55,"03")
text(s,1.25,0.48,9.2,0.75,"他们不是在念台词",size=36,color=CREAM,font=TITLE_FONT)
rule(s,0.82,1.35,1.1)
text(s,0.85,1.72,5.85,0.5,"千春盯着池田。池田怕她。",size=18,color=CREAM)
text(s,0.85,2.3,5.85,0.5,"小野对女儿要求很高，不轻易夸。",size=18,color=CREAM)
text(s,0.85,2.88,5.85,0.5,"樱井越被质疑，越要解释。",size=18,color=CREAM)
text(s,0.85,3.46,5.85,0.5,"茶终话少。幸子怕母亲失望。",size=18,color=CREAM)
text(s,0.85,4.04,5.85,0.5,"墨墨护短，但嘴上不认。",size=18,color=CREAM)
rule(s,0.85,4.85,5.4)
text(s,0.85,5.08,5.8,1.0,"狼会悍跳、倒钩、冲锋；预言家会起跳或藏；\n女巫和猎人会藏身份。",size=14,color=MUTED,spacing=1.45)
pic(s,"_ppt_assets/card_group.jpg",7.1,1.15,5.45,3.4)
text(s,7.1,4.75,5.45,0.8,"每句话都要接住上一位；\n投票和发言对得上。",size=15,color=CREAM,spacing=1.4)
footer(s,4)

# 05 how it runs
s=prs.slides.add_slide(blank); bg(s)
num(s,0.82,0.55,"04")
text(s,1.25,0.48,9.0,0.75,"它怎么跑起来",size=36,color=CREAM,font=TITLE_FONT)
rule(s,0.82,1.35,1.1)
blocks=[
 ("剧情骨架","入席、规则、抽身份、白天、投票、结局。\n由代码钉住，不跑飞。"),
 ("角色状态","身份、同伴、验人、药、怀疑谁。\n只有他自己知道。"),
 ("先想再说","站谁、点谁、投谁。\n先做计划，再开口。"),
 ("断网也能玩","本地台词兜底。\n联网时，角色更活。"),
]
bw=5.7; bh=1.75
coords=[(0.85,1.9),(6.75,1.9),(0.85,3.95),(6.75,3.95)]
for (title,desc),(x,y) in zip(blocks,coords):
    box=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(bw),Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb=DARK; box.line.color.rgb=RGBColor(0x4A,0x3A,0x2A); box.line.width=Pt(0.8); box.shadow.inherit=False
    text(s,x+0.3,y+0.22,bw-0.6,0.45,title,size=21,color=GOLD,font=TITLE_FONT)
    text(s,x+0.3,y+0.78,bw-0.6,0.75,desc,size=13,color=MUTED,spacing=1.35)
text(s,0.85,6.2,11.7,0.45,"角色不只换声线：每个人先看自己在局里的处境，再决定说什么。",size=13,color=MUTED)
footer(s,5)

# 06 now
s=prs.slides.add_slide(blank); bg(s)
num(s,0.82,0.55,"05")
text(s,1.25,0.48,8.5,0.75,"现在能玩到什么",size=36,color=CREAM,font=TITLE_FONT)
rule(s,0.82,1.35,1.1)
items=[
 ("完整序章 + 第一章","从洋馆日常走到月见茶会。"),
 ("AI 茶会能跑完一局","七个角色自己发言、投票、收尾。"),
 ("玩家能插话","点名、站边、投票，改变局面。"),
 ("Mac 桌宠入口","点「去洋馆」，直接进游戏。"),
]
y=1.85
for i,(h,d) in enumerate(items):
    text(s,0.85,y,0.55,0.4,f"{i+1}",size=12,color=GOLD,bold=True)
    text(s,1.35,y-0.04,6.0,0.4,h,size=17,color=CREAM)
    text(s,1.35,y+0.38,6.2,0.4,d,size=13,color=MUTED)
    y+=1.02
text(s,0.85,6.1,7.2,0.4,"下一步：更多角色、更长局、离线模型。",size=13,color=GOLD)
pic(s,"_ppt_assets/card_eye.jpg",7.55,1.45,4.95,3.7)
text(s,7.55,5.4,4.95,0.4,"现场可试玩",size=20,color=CREAM,font=TITLE_FONT,align=PP_ALIGN.CENTER)
text(s,7.55,5.9,4.95,0.3,"墨墨的洋馆",size=12,color=MUTED,align=PP_ALIGN.CENTER)
footer(s,6)

out="墨墨的洋馆-ACC2026-6页.pptx"
prs.save(out)
print("saved",out)
